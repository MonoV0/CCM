"""おみくじ候補の収集とExcel生成。Discordへの投稿は行わない。"""
import io

import discord
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

EMOJI_NAMES = {'Ni_logo2W', 'Ni_logo1B'}


def matches(message):
    return any(getattr(r.emoji, 'name', None) in EMOJI_NAMES and r.count > 0
               for r in message.reactions)


async def collect(channels, progress):
    rows, errors = [], []
    for channel in channels:
        try:
            async for message in channel.history(limit=None, oldest_first=True):
                if matches(message) and message.content.strip():
                    rows.append((message.id, message.content, message.author.name, message.jump_url))
        except discord.HTTPException:
            errors.append(str(channel.id))
        progress['done'] += 1
    rows.sort(key=lambda row: row[0])
    return rows, errors


def workbook_bytes(rows, errors, start):
    book = Workbook()
    sheet = book.active
    sheet.title = 'おみくじ'
    sheet.append(['番号', '文言', 'ユーザーID'])
    sources = book.create_sheet('出典')
    sources.append(['番号', '投稿URL'])
    for number, (_, text, username, url) in enumerate(rows, start):
        sheet.append([number, '', ''])
        for column, value in ((2, text), (3, username)):
            cell = sheet.cell(sheet.max_row, column)
            cell.value = ILLEGAL_CHARACTERS_RE.sub('', value)
            cell.data_type = 's'  # 投稿が「=」で始まっても数式として実行しない。
        sources.append([number, url])
    notes = book.create_sheet('使い方・取得結果')
    for text in ['番号・文言・ユーザーIDは自由に編集できます。',
                 'ユーザーIDは取得時点のユーザー名です。表示名ではありません。',
                 '対象: Ni_logo2W / Ni_logo1B。両方付いていても1投稿1件。',
                 '本文が空の投稿・スレッド内の投稿は対象外です。',
                 f'取得件数: {len(rows)}',
                 '取得失敗チャンネルID: ' + (', '.join(errors) or 'なし')]:
        notes.append([text])
    for ws in book:
        ws.freeze_panes = 'A2'
        ws.column_dimensions['A'].width = 12 if ws != notes else 90
        ws.column_dimensions['B'].width = 80
        ws.column_dimensions['C'].width = 28
        for row in ws:
            for cell in row:
                cell.font = Font(name='Arial', bold=cell.row == 1)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
    sheet.auto_filter.ref = sheet.dimensions
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()
